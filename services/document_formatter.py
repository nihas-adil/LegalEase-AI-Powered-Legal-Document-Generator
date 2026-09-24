from io import BytesIO
from html import escape
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import (
    WD_TABLE_ALIGNMENT,
    WD_CELL_VERTICAL_ALIGNMENT,
)

from fpdf import FPDF


# =========================================================
# PROJECT PATHS
# =========================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

# Combined Logo + LegalEase image
LOGO_PATH = PROJECT_ROOT / "Image" / "Logo and legalEasy.png"

# Windows Arial fonts
ARIAL_REGULAR = Path(r"C:\Windows\Fonts\arial.ttf")
ARIAL_BOLD = Path(r"C:\Windows\Fonts\arialbd.ttf")


# =========================================================
# FONT CHECK
# =========================================================

def check_pdf_fonts():
    if not ARIAL_REGULAR.exists():
        raise FileNotFoundError(
            f"Arial font not found at:\n{ARIAL_REGULAR}"
        )

    if not ARIAL_BOLD.exists():
        raise FileNotFoundError(
            f"Arial Bold font not found at:\n{ARIAL_BOLD}"
        )


# =========================================================
# LOGO CHECK
# =========================================================

def check_logo():
    if not LOGO_PATH.exists():
        raise FileNotFoundError(
            f"Logo not found at:\n{LOGO_PATH}"
        )


# =========================================================
# TXT
# =========================================================

def format_txt(text: str) -> bytes:
    return str(text).encode("utf-8")


# =========================================================
# HTML PREVIEW
# =========================================================

def format_html_preview(text: str) -> str:
    text = str(text)

    text = text.replace("**", "")
    text = text.replace("__", "")

    cleaned_lines = []

    for line in text.splitlines():
        line = line.strip()

        if line.startswith("#### "):
            line = line[5:].strip()
        elif line.startswith("### "):
            line = line[4:].strip()
        elif line.startswith("## "):
            line = line[3:].strip()
        elif line.startswith("# "):
            line = line[2:].strip()

        # Remove standalone LegalEase
        if line.lower() == "legalease":
            continue

        cleaned_lines.append(line)

    text = "\n".join(cleaned_lines)
    safe_text = escape(text)

    return (
        '<div class="legal-preview">'
        f"{safe_text}"
        "</div>"
    )


# =========================================================
# TEXT CLEANING
# =========================================================

def clean_document_text(text: str) -> str:
    text = str(text)

    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    text = text.replace("**", "")
    text = text.replace("__", "")

    text = text.replace("\u200b", "")
    text = text.replace("\ufeff", "")

    cleaned_lines = []

    for line in text.splitlines():
        line = line.strip()

        if line.startswith("#### "):
            line = line[5:].strip()
        elif line.startswith("### "):
            line = line[4:].strip()
        elif line.startswith("## "):
            line = line[3:].strip()
        elif line.startswith("# "):
            line = line[2:].strip()

        if line.lower() == "legalease":
            continue

        cleaned_lines.append(line)

    return "\n".join(cleaned_lines)


# =========================================================
# DOCX
# =========================================================

def format_docx(
    text: str,
    doc_type: str,
    parties: str,
    terms: str,
    effective_date: str,
) -> bytes:

    check_logo()

    document = Document()

    # -----------------------------------------------------
    # PAGE SETTINGS
    # -----------------------------------------------------

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # -----------------------------------------------------
    # DEFAULT FONT
    # -----------------------------------------------------

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    # -----------------------------------------------------
    # LOGO
    # -----------------------------------------------------

    logo_paragraph = document.add_paragraph()
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    logo_run = logo_paragraph.add_run()

    logo_run.add_picture(
        str(LOGO_PATH),
        width=Inches(1.8),
    )

    # -----------------------------------------------------
    # TITLE
    # -----------------------------------------------------

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title_paragraph.add_run(
        str(doc_type).upper()
    )

    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)

    # -----------------------------------------------------
    # EFFECTIVE DATE
    # -----------------------------------------------------

    date_paragraph = document.add_paragraph()
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    date_run = date_paragraph.add_run(
        f"Effective Date: {effective_date}"
    )

    date_run.font.name = "Times New Roman"
    date_run.font.size = Pt(11)

    document.add_paragraph()

    # -----------------------------------------------------
    # PARTIES
    # -----------------------------------------------------

    parties_heading = document.add_paragraph()

    parties_run = parties_heading.add_run("PARTIES")

    parties_run.bold = True
    parties_run.font.name = "Times New Roman"
    parties_run.font.size = Pt(13)

    parties_paragraph = document.add_paragraph()

    parties_run = parties_paragraph.add_run(
        str(parties)
    )

    parties_run.font.name = "Times New Roman"
    parties_run.font.size = Pt(11)

    # -----------------------------------------------------
    # TERMS TABLE
    # -----------------------------------------------------

    terms_heading = document.add_paragraph()

    terms_run = terms_heading.add_run(
        "TERMS AND CONDITIONS"
    )

    terms_run.bold = True
    terms_run.font.name = "Times New Roman"
    terms_run.font.size = Pt(13)

    terms_table = document.add_table(
        rows=1,
        cols=2,
    )

    terms_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    terms_table.style = "Table Grid"

    header_cells = terms_table.rows[0].cells

    header_cells[0].text = "No."
    header_cells[1].text = "Term / Condition"

    for cell in header_cells:

        cell.vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )

        for paragraph in cell.paragraphs:

            for run in paragraph.runs:

                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)

    # -----------------------------------------------------
    # ADD TERMS
    # -----------------------------------------------------

    term_items = []

    for item in str(terms).split(";"):

        item = item.strip()

        if item:
            term_items.append(item)

    if not term_items:
        term_items = [str(terms).strip()]

    for index, term in enumerate(
        term_items,
        start=1,
    ):

        row_cells = terms_table.add_row().cells

        row_cells[0].text = str(index)
        row_cells[1].text = term

        for cell in row_cells:

            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.TOP
            )

            for paragraph in cell.paragraphs:

                for run in paragraph.runs:

                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)

    # -----------------------------------------------------
    # GENERATED DOCUMENT CONTENT
    # -----------------------------------------------------

    document.add_paragraph()

    content_heading = document.add_paragraph()

    content_run = content_heading.add_run(
        "AGREEMENT"
    )

    content_run.bold = True
    content_run.font.name = "Times New Roman"
    content_run.font.size = Pt(13)

    cleaned_text = clean_document_text(text)

    headings = [
        "NON-DISCLOSURE AGREEMENT",
        "SERVICE AGREEMENT",
        "EMPLOYMENT AGREEMENT",
        "RENTAL AGREEMENT",
        "PARTNERSHIP AGREEMENT",
        "FREELANCE AGREEMENT",
        "FREELANCE WORK CONTRACT",
        "PARTIES",
        "TERMS AND CONDITIONS",
        "AGREEMENT",
        "SIGNATURES",
        "SERVICES AND DEADLINES",
        "SERVICES AND DELIVERY",
        "PAYMENT TERMS",
        "CONFIDENTIALITY",
        "TERM AND TERMINATION",
        "GOVERNING LAW",
        "ENTIRE AGREEMENT",
        "SEVERABILITY",
    ]

    for line in cleaned_text.splitlines():

        line = line.strip()

        if not line:

            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(6)

            continue

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.15

        if line.upper() in headings:

            run = paragraph.add_run(
                line.upper()
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)

        else:

            run = paragraph.add_run(line)

            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    # -----------------------------------------------------
    # SIGNATURE SECTION
    # -----------------------------------------------------

    signature_heading = document.add_paragraph()

    signature_run = signature_heading.add_run(
        "SIGNATURES"
    )

    signature_run.bold = True
    signature_run.font.name = "Times New Roman"
    signature_run.font.size = Pt(13)

    document.add_paragraph()

    signature_table = document.add_table(
        rows=4,
        cols=2,
    )

    signature_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    signature_table.style = "Table Grid"

    signature_table.cell(0, 0).text = "Party A"
    signature_table.cell(0, 1).text = "Party B"

    signature_table.cell(
        1,
        0,
    ).text = "Name: ______________________________"

    signature_table.cell(
        1,
        1,
    ).text = "Name: ______________________________"

    signature_table.cell(
        2,
        0,
    ).text = "Signature: __________________________"

    signature_table.cell(
        2,
        1,
    ).text = "Signature: __________________________"

    signature_table.cell(
        3,
        0,
    ).text = "Date: _______________________________"

    signature_table.cell(
        3,
        1,
    ).text = "Date: _______________________________"

    for row in signature_table.rows:

        for cell in row.cells:

            for paragraph in cell.paragraphs:

                for run in paragraph.runs:

                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)

    # -----------------------------------------------------
    # FOOTER
    # -----------------------------------------------------

    footer = section.footer.paragraphs[0]

    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer_run = footer.add_run(
        "LegalEase • AI-generated draft"
    )

    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(9)

    # -----------------------------------------------------
    # SAVE
    # -----------------------------------------------------

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


# =========================================================
# PDF CLASS
# =========================================================

class LegalPDF(FPDF):

    def __init__(
        self,
        *args,
        **kwargs,
    ):

        super().__init__(
            *args,
            **kwargs,
        )

        # -------------------------------------------------
        # REGISTER ARIAL FONTS
        # -------------------------------------------------

        self.add_font(
            "Arial",
            "",
            str(ARIAL_REGULAR),
        )

        self.add_font(
            "Arial",
            "B",
            str(ARIAL_BOLD),
        )

    # =====================================================
    # HEADER
    # =====================================================

    def header(self):

        if LOGO_PATH.exists():

            try:

                # Combined logo width
                logo_width = 45

                # Center on A4 page
                logo_x = (
                    self.w - logo_width
                ) / 2

                self.image(
                    str(LOGO_PATH),
                    x=logo_x,
                    y=8,
                    w=logo_width,
                )

            except Exception as e:

                print(
                    "Logo error:",
                    e,
                )

        # Start document content below logo
        self.set_y(30)

    # =====================================================
    # FOOTER
    # =====================================================

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Arial",
            "",
            9,
        )

        self.cell(
            0,
            10,
            f"LegalEase Inc. | Page {self.page_no()}",
            align="C",
        )


# =========================================================
# PDF
# =========================================================

def format_pdf(
    text: str,
    doc_type: str,
    effective_date: str,
) -> bytes:

    # -----------------------------------------------------
    # CHECK FILES
    # -----------------------------------------------------

    check_pdf_fonts()
    check_logo()

    # -----------------------------------------------------
    # CREATE PDF
    # -----------------------------------------------------

    pdf = LegalPDF(
        orientation="P",
        unit="mm",
        format="A4",
    )

    # -----------------------------------------------------
    # PAGE SETTINGS
    # -----------------------------------------------------

    pdf.set_margins(
        left=25,
        top=30,
        right=25,
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=25,
    )

    pdf.add_page()

    # -----------------------------------------------------
    # CLEAN TEXT
    # -----------------------------------------------------

    text = clean_document_text(text)

    # -----------------------------------------------------
    # FIX BULLETS
    # -----------------------------------------------------

    text = text.replace(
        "\uf0b7",
        "-",
    )

    text = text.replace(
        "",
        "-",
    )

    text = text.replace(
        "•",
        "-",
    )

    # -----------------------------------------------------
    # REMOVE DUPLICATES
    # -----------------------------------------------------

    cleaned_lines = []

    title_removed = False
    date_removed = False

    for raw_line in text.splitlines():

        line = raw_line.strip()

        # Empty line
        if not line:

            cleaned_lines.append("")

            continue

        # -------------------------------------------------
        # REMOVE DUPLICATE TITLE
        # -------------------------------------------------

        if (
            not title_removed
            and line.upper()
            == str(doc_type).upper()
        ):

            title_removed = True

            continue

        # -------------------------------------------------
        # REMOVE DUPLICATE DATE
        # -------------------------------------------------

        if (
            not date_removed
            and line.lower()
            == f"effective date: {effective_date}".lower()
        ):

            date_removed = True

            continue

        # -------------------------------------------------
        # REMOVE STANDALONE LEGALEASE
        # -------------------------------------------------

        if line.lower() == "legalease":

            continue

        # -------------------------------------------------
        # REMOVE MARKDOWN HEADINGS
        # -------------------------------------------------

        if line.startswith("#### "):

            line = line[5:].strip()

        elif line.startswith("### "):

            line = line[4:].strip()

        elif line.startswith("## "):

            line = line[3:].strip()

        elif line.startswith("# "):

            line = line[2:].strip()

        # -------------------------------------------------
        # REMOVE STANDALONE LEGALEASE AGAIN
        # -------------------------------------------------

        if line.lower() == "legalease":

            continue

        cleaned_lines.append(line)

    # -----------------------------------------------------
    # TITLE
    # -----------------------------------------------------

    pdf.set_text_color(
        0,
        0,
        0,
    )

    pdf.set_font(
        "Arial",
        "B",
        18,
    )

    pdf.multi_cell(
        0,
        9,
        str(doc_type).upper(),
        align="C",
    )

    pdf.ln(2)

    # -----------------------------------------------------
    # EFFECTIVE DATE
    # -----------------------------------------------------

    pdf.set_font(
        "Arial",
        "",
        11,
    )

    pdf.multi_cell(
        0,
        7,
        f"Effective Date: {effective_date}",
        align="C",
    )

    pdf.ln(8)

    # -----------------------------------------------------
    # DIVIDER
    # -----------------------------------------------------

    pdf.set_draw_color(
        150,
        150,
        150,
    )

    pdf.set_line_width(
        0.3,
    )

    pdf.line(
        pdf.l_margin,
        pdf.get_y(),
        pdf.w - pdf.r_margin,
        pdf.get_y(),
    )

    pdf.ln(8)

    # -----------------------------------------------------
    # HEADINGS
    # -----------------------------------------------------

    headings = [
        "PARTIES",
        "TERMS AND CONDITIONS",
        "AGREEMENT",
        "SIGNATURES",
        "FREELANCE WORK CONTRACT",
        "SERVICES AND DEADLINES",
        "SERVICES AND DELIVERY",
        "PAYMENT TERMS",
        "CONFIDENTIALITY",
        "TERM AND TERMINATION",
        "GOVERNING LAW",
        "ENTIRE AGREEMENT",
        "SEVERABILITY",
    ]

    # -----------------------------------------------------
    # CONTENT
    # -----------------------------------------------------

    for line in cleaned_lines:

        if not line:

            pdf.ln(4)

            continue

        heading_text = line.strip()

        is_numbered_heading = False

        # -------------------------------------------------
        # CHECK NUMBERED HEADINGS
        # -------------------------------------------------

        if "." in heading_text:

            first_part, remaining = (
                heading_text.split(
                    ".",
                    1,
                )
            )

            if (
                first_part.isdigit()
                and remaining.strip()
                and len(remaining.strip()) < 80
            ):

                possible_heading = remaining.strip()

                if possible_heading.isupper():

                    is_numbered_heading = True

        # -------------------------------------------------
        # HEADING
        # -------------------------------------------------

        if (
            line.upper() in headings
            or is_numbered_heading
        ):

            pdf.ln(3)

            pdf.set_font(
                "Arial",
                "B",
                12,
            )

            pdf.set_text_color(
                0,
                0,
                0,
            )

            pdf.multi_cell(
                0,
                7,
                line.upper(),
            )

            pdf.ln(2)

            continue

        # -------------------------------------------------
        # PARTY LINES
        # -------------------------------------------------

        if (
            line.startswith("Party A:")
            or line.startswith("Party B:")
            or line.startswith("SERVICE PROVIDER:")
            or line.startswith("CLIENT:")
        ):

            pdf.set_font(
                "Arial",
                "B",
                11,
            )

            pdf.multi_cell(
                0,
                8,
                line,
            )

            pdf.ln(2)

            continue

        # -------------------------------------------------
        # NORMAL CONTENT
        # -----------------------------------------------------

        pdf.set_font(
            "Arial",
            "",
            11,
        )

        pdf.set_text_color(
            0,
            0,
            0,
        )

        pdf.multi_cell(
            0,
            7,
            line,
            align="L",
        )

        pdf.ln(1)

    # -----------------------------------------------------
    # RETURN PDF
    # -----------------------------------------------------

    return bytes(
        pdf.output()
    )